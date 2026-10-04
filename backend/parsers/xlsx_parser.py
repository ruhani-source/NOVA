from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


def parse_xlsx(file_path: str) -> dict:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        # Read from bytes so openpyxl does not care about the fake extension.
        workbook = load_workbook(
            BytesIO(path.read_bytes()),
            data_only=True,
        )

        text_parts = []
        sheet_names = []

        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            text_parts.append(f"[Sheet: {sheet.title}]")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    text_parts.append(" | ".join(values))

        return {
            "filename": path.name,
            "file_type": "xlsx",
            "text": "\n".join(text_parts).strip(),
            "metadata": {
                "sheet_count": len(workbook.worksheets),
                "sheets": sheet_names,
            },
        }

    except (BadZipFile, ValueError, KeyError):
        # Some challenge files use .xlsx while actually containing plain text.
        try:
            text = path.read_text(
                encoding="utf-8",
                errors="strict",
            ).strip()

            return {
                "filename": path.name,
                "file_type": "txt",
                "text": text,
                "metadata": {
                    "original_extension": path.suffix.lower(),
                    "detected_type": "text",
                },
            }

        except (UnicodeDecodeError, OSError):
            raise ValueError(
                f"{path.name} is not a valid Excel or UTF-8 text file."
            )
