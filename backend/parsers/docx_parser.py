from pathlib import Path

from docx import Document


def parse_docx(file_path: str) -> dict:
    """
    Extract text and basic metadata from a DOCX file.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        A standardized dictionary containing the file information and text.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    document = Document(path)

    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    # Extract text from tables
    for table in document.tables:
        for row in table.rows:
            row_text = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if row_text:
                text_parts.append(" | ".join(row_text))

    full_text = "\n".join(text_parts).strip()

    return {
        "filename": path.name,
        "file_type": "docx",
        "text": full_text,
        "metadata": {
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
        },
    }